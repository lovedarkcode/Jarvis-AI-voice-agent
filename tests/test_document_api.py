"""Real document extraction and request-scoped RAG, without provider credentials."""
import unittest
import zipfile
from io import BytesIO
from unittest.mock import AsyncMock, patch

from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
from fastapi.testclient import TestClient

from server.provider_api import app
from server.documents import Attachment, document_context


def pdf_bytes(text='Invoice total is 431 dollars.'):
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
                             NameObject('/Subtype'): NameObject('/Type1'),
                             NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({
        NameObject('/F1'): writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(('BT /F1 12 Tf 20 250 Td (' + text + ') Tj ET').encode('ascii'))
    page[NameObject('/Contents')] = writer._add_object(stream)
    result = BytesIO()
    writer.write(result)
    return result.getvalue()


class DocumentApiTests(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)

    def upload(self, name, content):
        return self.client.post('/api/documents', files={'file': (name, content)})

    def test_pdf_with_short_text_layer_is_readable_and_citable(self):
        response = self.upload('invoice.pdf', pdf_bytes())
        self.assertEqual(response.status_code, 200, response.text)
        attachment = response.json()['attachment']
        self.assertEqual(attachment['name'], 'invoice.pdf')
        self.assertEqual(attachment['segments'][0]['label'], 'page 1')
        self.assertIn('431 dollars', attachment['segments'][0]['text'])
        self.assertEqual(response.headers['cache-control'], 'no-store')

    def test_general_reading_questions_keep_text_from_a_large_uploaded_pdf(self):
        response = self.upload('The_Hass_Is_The_System.pdf', pdf_bytes(
            'Designing useful systems requires clear ownership and feedback loops. ' * 350))
        self.assertEqual(response.status_code, 200, response.text)
        self.assertGreater(response.json()['characters'], 12000)
        attachment = Attachment(**response.json()['attachment'])
        for question in ('Can you check right now?', 'Can you see the PDF Ram?',
                         'Read this PDF', 'What is this document about?'):
            with self.subTest(question=question):
                context = document_context(attachment, question, question)
                self.assertIn('ATTACHMENT AVAILABLE', context)
                self.assertIn('The_Hass_Is_The_System.pdf', context)
                self.assertIn('Designing useful systems requires clear ownership', context)
                self.assertIn('page 1', context)
                self.assertNotIn('Summarise what this document', context)
                self.assertLess(len(context), 24000)

    def test_missing_topic_still_reports_attachment_available_without_inventing_an_answer(self):
        attachment = Attachment(name='design.pdf', kind='pdf', segments=[
            {'label': 'page 1', 'text': 'Clear ownership and feedback loops help design systems. ' * 400}])
        context = document_context(attachment, 'What was the dinosaur population?')
        self.assertIn('ATTACHMENT AVAILABLE', context)
        self.assertIn('No focused passage matched', context)
        self.assertIn('Representative excerpts', context)
        self.assertIn('using ONLY the text above', context)
        self.assertIn('Do not claim the upload failed', context)

    def test_provider_prompts_use_current_pdf_even_after_previous_missing_pdf_replies(self):
        attachment = {'name': 'design.pdf', 'kind': 'pdf', 'segments': [
            {'label': 'page 1', 'text': 'Designing systems requires clear ownership. ' * 400}]}
        for provider in ('openai', 'claude', 'sarvam', 'gemini'):
            with self.subTest(provider=provider), patch('server.provider_api.provider_call', new=AsyncMock(
                return_value={'output': [{'content': [{'type': 'output_text', 'text': 'Ready'}]}],
                              'content': [{'type': 'text', 'text': 'Ready'}],
                              'choices': [{'message': {'content': 'Ready'}}],
                              'candidates': [{'content': {'parts': [{'text': 'Ready'}]}}]})) as call:
                response = self.client.post('/api/chat', json={
                    'provider': provider, 'key': 'a' * 32,
                    'messages': [{'role': 'user', 'content': 'Read my PDF'},
                                 {'role': 'assistant', 'content': 'No PDF was uploaded. Please upload again.'},
                                 {'role': 'user', 'content': 'Can you see the PDF Ram?'}],
                    'attachment': attachment})
                self.assertEqual(response.status_code, 200, response.text)
                payload = call.call_args.kwargs['json']
                prompt = (payload['instructions'] if provider == 'openai' else payload['system'] if provider == 'claude'
                          else payload['messages'][0]['content'] if provider == 'sarvam'
                          else payload['systemInstruction']['parts'][0]['text'])
                self.assertIn('A document is currently attached', prompt)
                self.assertIn('takes precedence over earlier conversation messages', prompt)
                self.assertIn('ATTACHMENT AVAILABLE', prompt)
                self.assertIn('Designing systems requires clear ownership.', prompt)
                self.assertIn('page 1', prompt)

    def test_word_preserves_paragraphs_and_tables(self):
        document = Document()
        document.add_heading('Warranty', level=1)
        document.add_paragraph('The warranty lasts 18 months.')
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = 'Plan'
        table.cell(0, 1).text = 'Cost'
        table.cell(1, 0).text = 'Premium'
        table.cell(1, 1).text = '49'
        buffer = BytesIO()
        document.save(buffer)
        response = self.upload('policy.docx', buffer.getvalue())
        self.assertEqual(response.status_code, 200, response.text)
        text = '\n'.join(s['text'] for s in response.json()['attachment']['segments'])
        for expected in ('18 months', 'Premium', '49'):
            self.assertIn(expected, text)

    def test_excel_preserves_sheet_and_headers_across_row_groups(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Sales'
        sheet.append(['Product', 'Revenue'])
        for i in range(250):
            sheet.append(['Widget ' + str(i), i * 3])
        buffer = BytesIO()
        workbook.save(buffer)
        response = self.upload('sales.xlsx', buffer.getvalue())
        self.assertEqual(response.status_code, 200, response.text)
        segments = response.json()['attachment']['segments']
        self.assertEqual(len(segments), 2)
        self.assertIn("sheet 'Sales'", segments[1]['label'])
        self.assertIn('Product | Revenue', segments[1]['text'])
        self.assertIn('Widget 249 | 747', segments[1]['text'])

    def test_all_chat_providers_receive_retrieved_passages_far_beyond_old_limit(self):
        content = ('Routine equipment maintenance schedule.\n' * 1700
                   + '\nZircon access code is 9274.\n')
        upload = self.upload('manual.txt', content.encode())
        self.assertEqual(upload.status_code, 200, upload.text)
        for provider in ('openai', 'claude', 'sarvam', 'gemini'):
            with self.subTest(provider=provider), patch('server.provider_api.provider_call', new=AsyncMock(
                return_value={'output': [{'content': [{'type': 'output_text', 'text': '9274'}]}],
                              'content': [{'type': 'text', 'text': '9274'}],
                              'choices': [{'message': {'content': '9274'}}],
                              'candidates': [{'content': {'parts': [{'text': '9274'}]}}]})) as call:
                response = self.client.post('/api/chat', json={
                    'provider': provider, 'key': 'AIza' + 'a' * 32 if provider == 'gemini' else 'sk-' + 'a' * 32,
                    'messages': [{'role': 'user', 'content': 'What is the Zircon access code?'}],
                    'attachment': upload.json()['attachment']})
                self.assertEqual(response.status_code, 200, response.text)
                payload = call.call_args.kwargs['json']
                prompt = (payload['instructions'] if provider == 'openai' else payload['system'] if provider == 'claude'
                          else payload['messages'][0]['content'] if provider == 'sarvam'
                          else payload['systemInstruction']['parts'][0]['text'])
                self.assertIn('Zircon access code is 9274', prompt)
                self.assertIn('manual.txt', prompt)
                self.assertLess(len(prompt), 25000)

    def test_uploaded_text_is_not_implicitly_shared_with_other_requests(self):
        self.upload('private.txt', b'Private password is pineapple.')
        with patch('server.provider_api.provider_call', new=AsyncMock(return_value={
            'choices': [{'message': {'content': 'Hi'}}]})) as call:
            self.client.post('/api/chat', json={'provider': 'sarvam', 'key': 'a' * 32,
                             'messages': [{'role': 'user', 'content': 'Hi'}]})
        self.assertNotIn('pineapple', str(call.call_args.kwargs))

    def test_summary_samples_the_document_and_followup_uses_retrieval(self):
        attachment = Attachment(name='manual.txt', kind='text', segments=[
            {'label': 'Maintenance', 'text': 'Routine equipment maintenance schedule.\n' * 1700},
            {'label': 'Zircon', 'text': 'Zircon access code is 9274.'}])
        overview = document_context(attachment, 'Summarize the document')
        self.assertIn('SAMPLE', overview)
        followup = document_context(attachment, 'Summarize the document\nWhat is the Zircon code?', 'What is the Zircon code?')
        self.assertIn('Zircon access code is 9274.', followup)
        self.assertIn('most relevant passages', followup)

    def test_missing_empty_corrupt_and_unsupported_files_have_clear_errors(self):
        self.assertEqual(self.client.post('/api/documents').status_code, 400)
        for name, content in [('empty.txt', b''), ('broken.pdf', b'not a PDF'),
                              ('broken.docx', b'not a ZIP'), ('legacy.doc', b'old Word'),
                              ('legacy.xls', b'old Excel'), ('empty.pdf', pdf_bytes(''))]:
            with self.subTest(name=name):
                response = self.upload(name, content)
                self.assertEqual(response.status_code, 400, response.text)
                self.assertIn('error', response.json())

    def test_upload_and_extracted_text_limits(self):
        self.assertEqual(self.upload('large.txt', b'x' * 3_500_001).status_code, 413)
        self.assertEqual(self.upload('long.txt', b'x' * 400_001).status_code, 413)
        self.assertEqual(self.upload('request.txt', b'x' * 4_000_001).status_code, 413)

    def test_office_expansion_is_bounded_before_parsing(self):
        buffer = BytesIO()
        with zipfile.ZipFile(buffer, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr('word/document.xml', b'x' * 16_000_001)
        self.assertEqual(self.upload('oversized.docx', buffer.getvalue()).status_code, 413)

    def test_client_supplied_segments_are_bounded_and_validation_redacted(self):
        response = self.client.post('/api/chat', json={'provider': 'sarvam', 'key': 'sensitive-key-123456',
            'messages': [{'role': 'user', 'content': 'Hi'}],
            'attachment': {'name': 'large.txt', 'kind': 'text',
                           'segments': [{'label': 'section', 'text': 'private' * 70000}]}})
        self.assertEqual(response.status_code, 422)
        self.assertNotIn('sensitive-key', response.text)
        self.assertNotIn('private', response.text)


if __name__ == '__main__':
    unittest.main()
